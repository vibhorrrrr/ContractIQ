"""Parsing service — extract text from PDF and DOCX files."""

import logging
from pathlib import Path

import pdfplumber
import fitz  # PyMuPDF
from docx import Document

logger = logging.getLogger(__name__)


class ParsingService:
    """Extracts raw text and page counts from uploaded contract files."""

    async def parse(self, file_path: str) -> tuple[str, int]:
        """Parse a document and return (full_text, page_count).

        Raises ValueError if the file cannot be parsed.
        """
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext == ".pdf":
            return await self._parse_pdf(path)
        elif ext == ".docx":
            return await self._parse_docx(path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

    async def _parse_pdf(self, path: Path) -> tuple[str, int]:
        """Extract text from PDF using pdfplumber, falling back to PyMuPDF."""
        # Try pdfplumber first
        try:
            return self._extract_with_pdfplumber(path)
        except Exception as e:
            logger.warning("pdfplumber failed for %s: %s. Trying PyMuPDF.", path, e)

        # Fallback to PyMuPDF
        try:
            return self._extract_with_pymupdf(path)
        except Exception as e:
            logger.error("PyMuPDF also failed for %s: %s", path, e)
            raise ValueError("File could not be read. Please re-upload.") from e

    def _extract_with_pdfplumber(self, path: Path) -> tuple[str, int]:
        """Extract text using pdfplumber (best for complex layouts)."""
        pages_text: list[str] = []
        with pdfplumber.open(path) as pdf:
            if pdf.is_encrypted:
                raise ValueError("Password-protected files are not supported.")

            for page in pdf.pages:
                text = page.extract_text() or ""
                pages_text.append(text)

        full_text = "\n\n".join(pages_text)
        page_count = len(pages_text)

        if not full_text.strip():
            raise ValueError(
                "This PDF appears to be a scanned image. "
                "Please upload a text-based PDF."
            )

        return full_text, page_count

    def _extract_with_pymupdf(self, path: Path) -> tuple[str, int]:
        """Extract text using PyMuPDF as fallback."""
        doc = fitz.open(str(path))
        pages_text: list[str] = []

        for page in doc:
            text = page.get_text()
            pages_text.append(text)

        doc.close()
        full_text = "\n\n".join(pages_text)
        page_count = len(pages_text)

        if not full_text.strip():
            raise ValueError(
                "This PDF appears to be a scanned image. "
                "Please upload a text-based PDF."
            )

        return full_text, page_count

    async def _parse_docx(self, path: Path) -> tuple[str, int]:
        """Extract text from DOCX using python-docx."""
        try:
            doc = Document(str(path))
        except Exception as e:
            raise ValueError("File could not be read. Please re-upload.") from e

        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        paragraphs.append(cell.text)

        full_text = "\n\n".join(paragraphs)

        if not full_text.strip():
            raise ValueError("Document appears to be empty.")

        # DOCX doesn't have a direct page count concept; estimate from content
        estimated_pages = max(1, len(full_text) // 3000)
        return full_text, estimated_pages
